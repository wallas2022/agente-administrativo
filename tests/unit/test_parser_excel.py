import io
from datetime import date

from openpyxl import Workbook

from parsers.excel import leer_libro_contable


def _libro_de_prueba(*, fila_total_con_formula: bool) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Partidas"
    ws.append(["Cuenta", "Descripcion", "Fecha", "Debe", "Haber", "Moneda"])
    ws.append(["1010", "Cobro cliente A", date(2026, 1, 5), 1000.00, 0.00, "Q"])
    ws.append(["4010", "Ingreso servicio A", date(2026, 1, 5), 0.00, 1000.00, "Q"])
    ws.append(["1010", "Cobro cliente B", date(2026, 1, 10), 500.00, 0.00, "Q"])
    ws.append(["4010", "Ingreso servicio B", date(2026, 1, 10), 0.00, 500.00, "Q"])

    fila_total = 6
    if fila_total_con_formula:
        ws.cell(row=fila_total, column=1, value="TOTAL")
        ws.cell(row=fila_total, column=4, value="=SUM(D2:D5)")
        ws.cell(row=fila_total, column=5, value="=SUM(E2:E5)")
    else:
        ws.cell(row=fila_total, column=1, value="TOTAL")
        ws.cell(row=fila_total, column=4, value=1500.00)  # valor fijo, no fórmula
        ws.cell(row=fila_total, column=5, value=1500.00)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_leer_libro_extrae_las_partidas_con_ubicacion() -> None:
    libro = leer_libro_contable(io.BytesIO(_libro_de_prueba(fila_total_con_formula=True)))

    assert len(libro.partidas) == 4
    primera = libro.partidas[0]
    assert primera.hoja == "Partidas"
    assert primera.fila == 2
    assert primera.cuenta == "1010"
    assert primera.debe == 1000.00
    assert primera.moneda == "Q"


def test_leer_libro_detecta_si_la_fila_total_es_formula() -> None:
    libro_con_formula = leer_libro_contable(
        io.BytesIO(_libro_de_prueba(fila_total_con_formula=True))
    )
    libro_sin_formula = leer_libro_contable(
        io.BytesIO(_libro_de_prueba(fila_total_con_formula=False))
    )

    assert libro_con_formula.fila_total is not None
    assert libro_con_formula.fila_total.debe_es_formula is True
    assert libro_con_formula.fila_total.haber_es_formula is True

    assert libro_sin_formula.fila_total is not None
    assert libro_sin_formula.fila_total.debe_es_formula is False
    assert libro_sin_formula.fila_total.haber_es_formula is False
