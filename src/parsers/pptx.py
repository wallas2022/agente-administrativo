"""Extractor de texto de PowerPoint (RF-10, CU-05): título, viñetas y notas
del orador por diapositiva. Solo lectura -- la revisión ortográfica vive en
`src/ortografia`.
"""

from __future__ import annotations

import io

from pptx import Presentation

from parsers.segmentos import SegmentoTexto


def leer_texto_pptx(archivo: io.BytesIO | str) -> list[SegmentoTexto]:
    if hasattr(archivo, "seek"):
        archivo.seek(0)
    presentacion = Presentation(archivo)

    segmentos: list[SegmentoTexto] = []
    for si, diapositiva in enumerate(presentacion.slides, start=1):
        forma_titulo = diapositiva.shapes.title
        # `shapes.title` construye un wrapper nuevo en cada acceso -- `is`
        # nunca coincide con la forma del bucle aunque sea la misma diapositiva
        # subyacente; shape_id sí es estable.
        id_forma_titulo = forma_titulo.shape_id if forma_titulo is not None else None
        contador_vineta = 0
        for forma in diapositiva.shapes:
            if not forma.has_text_frame:
                continue
            es_titulo = forma.shape_id == id_forma_titulo
            for parrafo in forma.text_frame.paragraphs:
                texto = "".join(corrida.text for corrida in parrafo.runs)
                if not texto.strip():
                    continue
                if es_titulo:
                    ubicacion = f"Diapositiva {si}, título"
                else:
                    contador_vineta += 1
                    ubicacion = f"Diapositiva {si}, viñeta {contador_vineta}"
                segmentos.append(SegmentoTexto(texto=texto, ubicacion=ubicacion))

        if diapositiva.has_notes_slide:
            notas = diapositiva.notes_slide.notes_text_frame.text
            if notas.strip():
                segmentos.append(
                    SegmentoTexto(texto=notas, ubicacion=f"Diapositiva {si}, notas")
                )

    return segmentos
