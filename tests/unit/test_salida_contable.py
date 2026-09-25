import io
import json

from openpyxl import Workbook, load_workbook

from validadores.contable.reglas import HallazgoDetectado
from validadores.contable.salida import hallazgos_a_json, marcar_celdas_en_libro


def _libro_simple_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Partidas"
    ws.append(["Cuenta", "Descripcion", "Fecha", "Debe", "Haber", "Moneda"])
    ws.append(["1010", "Cobro", "2026-01-05", 100.0, 0.0, "Q"])
    ws.append(["9999", "Cuenta mala", "2026-01-05", 0.0, 100.0, "Q"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _hallazgo(severidad: str = "alta") -> HallazgoDetectado:
    return HallazgoDetectado(
        regla_codigo="RN-02",
        severidad=severidad,
        hoja="Partidas",
        fila=3,
        ubicacion="Partidas!A3",
        descripcion="La cuenta '9999' no existe en el catálogo",
    )


def test_marcar_celdas_colorea_la_celda_del_hallazgo_y_agrega_comentario() -> None:
    original = io.BytesIO(_libro_simple_bytes())

    resultado = marcar_celdas_en_libro(original, [_hallazgo()])

    wb = load_workbook(io.BytesIO(resultado))
    celda = wb["Partidas"]["A3"]
    assert celda.fill.fgColor.rgb in ("00FF0000", "FFFF0000")  # rojo: severidad alta
    assert celda.comment is not None
    assert "RN-02" in celda.comment.text


def test_marcar_celdas_no_modifica_el_libro_original() -> None:
    original_bytes = _libro_simple_bytes()
    original = io.BytesIO(original_bytes)

    marcar_celdas_en_libro(original, [_hallazgo()])

    assert original.getvalue() == original_bytes


def test_hallazgos_a_json_serializa_campos_clave() -> None:
    salida = hallazgos_a_json([_hallazgo()])
    datos = json.loads(salida)

    assert len(datos) == 1
    assert datos[0]["regla_codigo"] == "RN-02"
    assert datos[0]["ubicacion"] == "Partidas!A3"
    assert datos[0]["severidad"] == "alta"
