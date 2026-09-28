import io

from pptx import Presentation
from pptx.util import Inches

from parsers.correcciones import CorreccionAplicable
from parsers.pptx import aplicar_correcciones_pptx, leer_texto_pptx


def _pptx_de_prueba() -> bytes:
    presentacion = Presentation()
    diapositiva = presentacion.slides.add_slide(presentacion.slide_layouts[1])
    diapositiva.shapes.title.text = "Título de la diapositiva"
    marcador = diapositiva.placeholders[1]
    marcador.text_frame.text = "Primera viñeta"
    marcador.text_frame.add_paragraph().text = "Segunda viñeta"

    cuadro = diapositiva.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(1))
    cuadro.text_frame.text = "Cuadro de texto adicional"

    notas = diapositiva.notes_slide.notes_text_frame
    notas.text = "Notas del orador para esta diapositiva"

    buffer = io.BytesIO()
    presentacion.save(buffer)
    return buffer.getvalue()


def test_extrae_titulo_de_la_diapositiva() -> None:
    segmentos = leer_texto_pptx(io.BytesIO(_pptx_de_prueba()))

    titulo = [s for s in segmentos if s.ubicacion == "Diapositiva 1, título"]
    assert len(titulo) == 1
    assert titulo[0].texto == "Título de la diapositiva"


def test_extrae_vinetas_numeradas_sin_colisionar_entre_formas() -> None:
    segmentos = leer_texto_pptx(io.BytesIO(_pptx_de_prueba()))

    vinetas = {s.ubicacion: s.texto for s in segmentos if "viñeta" in s.ubicacion}
    assert vinetas["Diapositiva 1, viñeta 1"] == "Primera viñeta"
    assert vinetas["Diapositiva 1, viñeta 2"] == "Segunda viñeta"
    # El cuadro de texto adicional es otra forma no-título: sigue la misma
    # numeración de viñeta que las del marcador, sin reiniciar en 1.
    assert vinetas["Diapositiva 1, viñeta 3"] == "Cuadro de texto adicional"


def test_extrae_notas_del_orador() -> None:
    segmentos = leer_texto_pptx(io.BytesIO(_pptx_de_prueba()))

    notas = [s for s in segmentos if s.ubicacion == "Diapositiva 1, notas"]
    assert len(notas) == 1
    assert notas[0].texto == "Notas del orador para esta diapositiva"


def test_aplicar_correcciones_corrige_titulo_vineta_y_notas() -> None:
    correcciones = [
        CorreccionAplicable("Diapositiva 1, título", "Título de la diapositiva", "Título correcto"),
        CorreccionAplicable("Diapositiva 1, viñeta 2", "Segunda viñeta", "Segunda viñeta correcta"),
        CorreccionAplicable(
            "Diapositiva 1, notas",
            "Notas del orador para esta diapositiva",
            "Notas corregidas",
        ),
    ]

    resultado = aplicar_correcciones_pptx(io.BytesIO(_pptx_de_prueba()), correcciones)

    assert resultado.no_aplicadas == []
    segmentos = leer_texto_pptx(io.BytesIO(resultado.contenido))
    textos = {s.texto for s in segmentos}
    assert "Título correcto" in textos
    assert "Segunda viñeta correcta" in textos
    assert "Notas corregidas" in textos
    # La primera viñeta y el cuadro adicional no se tocaron.
    assert "Primera viñeta" in textos
    assert "Cuadro de texto adicional" in textos


def test_correccion_con_ubicacion_inexistente_queda_en_no_aplicadas() -> None:
    correccion = CorreccionAplicable("Diapositiva 9, título", "algo", "otra cosa")

    resultado = aplicar_correcciones_pptx(io.BytesIO(_pptx_de_prueba()), [correccion])

    assert resultado.no_aplicadas == [correccion]
