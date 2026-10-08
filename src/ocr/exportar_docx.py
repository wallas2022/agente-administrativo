"""Exporta el texto reconocido por OCR a .docx con las palabras dudosas
resaltadas (CU-06, Bloque 3): rojo para nivel "dudosa", amarillo para
"revisar" -- mismo vocabulario de `ocr.calidad`. Nunca modifica el texto en
sí, solo el formato visual (RN-06: la descarga muestra exactamente lo que
Tesseract reconoció).
"""

from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

_PATRON_TOKEN = re.compile(r"\S+|\s+")

_COLOR_POR_NIVEL = {
    "dudosa": WD_COLOR_INDEX.RED,
    "revisar": WD_COLOR_INDEX.YELLOW,
}


def construir_docx_resaltado(texto: str, *, nivel_por_palabra: dict[str, str]) -> bytes:
    """`nivel_por_palabra`: texto exacto de la palabra (tal como la reconoció
    Tesseract) -> nivel ("dudosa"/"revisar"). Coincidencia exacta de texto,
    no por posición -- una misma palabra dudosa en varias ocurrencias se
    resalta en todas (simplificación razonable para esta fase, documentada
    en docs/04-pruebas/resultados/local-cu06-bloque3.md)."""
    documento = Document()
    for linea in texto.split("\n"):
        parrafo = documento.add_paragraph()
        for token in _PATRON_TOKEN.findall(linea):
            run = parrafo.add_run(token)
            nivel = nivel_por_palabra.get(token.strip())
            color = _COLOR_POR_NIVEL.get(nivel) if nivel else None
            if color is not None:
                run.font.highlight_color = color

    buffer = io.BytesIO()
    documento.save(buffer)
    return buffer.getvalue()
