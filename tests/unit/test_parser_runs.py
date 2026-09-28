from docx import Document

from parsers.runs import reemplazar_texto_en_parrafo


def _parrafo_con_corridas(*fragmentos: tuple[str, bool]) -> object:
    """Crea un párrafo cuyas corridas tienen exactamente los textos dados
    (bold=True marca una corrida en negrita)."""
    documento = Document()
    parrafo = documento.add_paragraph()
    for texto, negrita in fragmentos:
        corrida = parrafo.add_run(texto)
        corrida.bold = negrita
    return parrafo


def test_reemplazo_dentro_de_una_sola_corrida_preserva_las_demas() -> None:
    parrafo = _parrafo_con_corridas(("Antes ", False), ("aprovado", False), (" después", False))

    aplicado = reemplazar_texto_en_parrafo(parrafo, "aprovado", "aprobado")

    assert aplicado is True
    assert "".join(c.text for c in parrafo.runs) == "Antes aprobado después"
    assert len(parrafo.runs) == 3


def test_reemplazo_dentro_de_corrida_en_negrita_conserva_el_formato() -> None:
    parrafo = _parrafo_con_corridas(
        ("Aplica a todo. ", False), ("Ningun asiento", True), (" más.", False)
    )

    aplicado = reemplazar_texto_en_parrafo(parrafo, "Ningun", "Ningún")

    assert aplicado is True
    corridas = list(parrafo.runs)
    assert corridas[1].text == "Ningún asiento"
    assert corridas[1].bold is True
    assert corridas[0].text == "Aplica a todo. "
    assert corridas[0].bold is not True


def test_reemplazo_que_cruza_el_limite_entre_dos_corridas() -> None:
    # "Ningun" queda partido entre dos corridas: "Ning" (normal) + "un asiento" (negrita).
    parrafo = _parrafo_con_corridas(("Ning", False), ("un asiento", True))

    aplicado = reemplazar_texto_en_parrafo(parrafo, "Ningun", "Ningún")

    assert aplicado is True
    assert "".join(c.text for c in parrafo.runs) == "Ningún asiento"
    # El texto nuevo completo queda en la primera corrida tocada (su formato).
    assert parrafo.runs[0].text == "Ningún"
    assert parrafo.runs[0].bold is not True
    assert parrafo.runs[1].text == " asiento"
    assert parrafo.runs[1].bold is True


def test_texto_no_encontrado_no_modifica_el_parrafo() -> None:
    parrafo = _parrafo_con_corridas(("Todo correcto aquí.", False))

    aplicado = reemplazar_texto_en_parrafo(parrafo, "inexistente", "reemplazo")

    assert aplicado is False
    assert parrafo.runs[0].text == "Todo correcto aquí."


def test_solo_reemplaza_la_primera_aparicion() -> None:
    parrafo = _parrafo_con_corridas(("gato gato gato", False))

    reemplazar_texto_en_parrafo(parrafo, "gato", "perro")

    assert parrafo.runs[0].text == "perro gato gato"


def test_texto_original_corto_no_colisiona_con_la_misma_letra_dentro_de_otra_palabra() -> None:
    """Bug real encontrado verificando el Bloque O4 en vivo: corregir la 'a'
    suelta de 'procedimiento a sido' -> 'ha sido' encontraba, en cambio, la
    'a' de 'Establecer' (la primera aparición por orden de texto), dejando
    'Esthablecer'. localizar_ocurrencia debe preferir el límite de palabra."""
    parrafo = _parrafo_con_corridas(
        ("Establecer los pasos. Este procedimiento a sido aprobado.", False)
    )

    reemplazar_texto_en_parrafo(parrafo, "a", "ha")

    assert parrafo.runs[0].text == "Establecer los pasos. Este procedimiento ha sido aprobado."
