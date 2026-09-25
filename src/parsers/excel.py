"""Parser de Excel contable (RF-06, CU-01). Extrae partidas y detecta si la
fila de totales usa fórmula o un valor fijo (RN: "fórmula reemplazada por
valor"). Sin lógica de validación aquí — solo lectura y normalización; las
reglas de negocio viven en `validadores.contable`.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook

CUENTA_TOTAL = "TOTAL"


@dataclass(frozen=True)
class PartidaContable:
    hoja: str
    fila: int
    cuenta: str
    descripcion: str
    fecha: date | None
    debe: float
    haber: float
    moneda: str


@dataclass(frozen=True)
class FilaTotal:
    hoja: str
    fila: int
    debe_declarado: float
    haber_declarado: float
    debe_es_formula: bool
    haber_es_formula: bool


@dataclass(frozen=True)
class LibroContable:
    partidas: list[PartidaContable] = field(default_factory=list)
    fila_total: FilaTotal | None = None


def _a_numero(valor: object) -> float:
    if valor is None:
        return 0.0
    if isinstance(valor, int | float):
        return float(valor)
    return 0.0


def _a_fecha(valor: object) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return None


def leer_libro_contable(archivo: io.BytesIO | str, hoja: str | None = None) -> LibroContable:
    """Lee un Excel contable con columnas Cuenta|Descripcion|Fecha|Debe|Haber|Moneda
    y una fila opcional "TOTAL" al final. Se hacen dos lecturas (con y sin
    `data_only`) para poder distinguir una fórmula de un valor fijo (RNF-03:
    el parser no calcula nada, solo reporta lo que hay en la hoja).
    """
    if hasattr(archivo, "seek"):
        archivo.seek(0)
    wb_formulas = load_workbook(archivo, data_only=False)
    if hasattr(archivo, "seek"):
        archivo.seek(0)
    wb_valores = load_workbook(archivo, data_only=True)

    nombre_hoja = hoja or wb_formulas.sheetnames[0]
    ws_formulas = wb_formulas[nombre_hoja]
    ws_valores = wb_valores[nombre_hoja]

    partidas: list[PartidaContable] = []
    fila_total: FilaTotal | None = None

    for num_fila in range(2, ws_formulas.max_row + 1):
        cuenta = ws_valores.cell(row=num_fila, column=1).value
        if cuenta is None:
            continue
        cuenta = str(cuenta).strip()
        if not cuenta:
            continue

        if cuenta.upper() == CUENTA_TOTAL:
            celda_debe_f = ws_formulas.cell(row=num_fila, column=4)
            celda_haber_f = ws_formulas.cell(row=num_fila, column=5)
            fila_total = FilaTotal(
                hoja=nombre_hoja,
                fila=num_fila,
                debe_declarado=_a_numero(ws_valores.cell(row=num_fila, column=4).value),
                haber_declarado=_a_numero(ws_valores.cell(row=num_fila, column=5).value),
                debe_es_formula=celda_debe_f.data_type == "f",
                haber_es_formula=celda_haber_f.data_type == "f",
            )
            continue

        partidas.append(
            PartidaContable(
                hoja=nombre_hoja,
                fila=num_fila,
                cuenta=cuenta,
                descripcion=str(ws_valores.cell(row=num_fila, column=2).value or ""),
                fecha=_a_fecha(ws_valores.cell(row=num_fila, column=3).value),
                debe=_a_numero(ws_valores.cell(row=num_fila, column=4).value),
                haber=_a_numero(ws_valores.cell(row=num_fila, column=5).value),
                moneda=str(ws_valores.cell(row=num_fila, column=6).value or "").strip().upper(),
            )
        )

    return LibroContable(partidas=partidas, fila_total=fila_total)
