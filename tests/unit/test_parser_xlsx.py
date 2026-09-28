import io

from openpyxl import Workbook

from parsers.xlsx import leer_texto_xlsx


def _libro_de_prueba() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Descripciones"
    ws.append(["Código", "Descripción"])
    ws.append(["1020", "Bancos, cuentas monetarias"])
    ws.append([100, "Gastos de papelería"])  # A3 numérico: no debe extraerse
    ws.cell(row=4, column=1, value="=A2")  # fórmula que evalúa a texto: excluir igual
    ws.cell(row=4, column=1).data_type = "f"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_extrae_celdas_de_texto_con_su_coordenada() -> None:
    segmentos = leer_texto_xlsx(io.BytesIO(_libro_de_prueba()))

    por_ubicacion = {s.ubicacion: s.texto for s in segmentos}
    assert por_ubicacion["Descripciones!A1"] == "Código"
    assert por_ubicacion["Descripciones!B2"] == "Bancos, cuentas monetarias"


def test_no_extrae_celdas_numericas() -> None:
    segmentos = leer_texto_xlsx(io.BytesIO(_libro_de_prueba()))

    assert not any(s.ubicacion == "Descripciones!A3" for s in segmentos)


def test_no_extrae_celdas_con_formula() -> None:
    segmentos = leer_texto_xlsx(io.BytesIO(_libro_de_prueba()))

    assert not any(s.ubicacion == "Descripciones!A4" for s in segmentos)
