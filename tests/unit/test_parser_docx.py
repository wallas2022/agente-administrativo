import io

from docx import Document

from parsers.correcciones import CorreccionAplicable
from parsers.docx import aplicar_correcciones_docx, construir_docx_desde_segmentos, leer_texto_docx
from parsers.segmentos import SegmentoTexto


def _docx_de_prueba() -> bytes:
    documento = Document()
    documento.add_paragraph("Título del documento")
    parrafo = documento.add_paragraph("Texto normal y ")
    parrafo.add_run("texto en negrita").bold = True

    tabla = documento.add_table(rows=2, cols=2)
    tabla.rows[0].cells[0].text = "Encabezado columna 1"
    tabla.rows[0].cells[1].text = "Encabezado columna 2"
    tabla.rows[1].cells[0].text = "Valor A"
    tabla.rows[1].cells[1].text = "Valor B"

    seccion = documento.sections[0]
    seccion.header.paragraphs[0].text = "Encabezado de página"
    seccion.footer.paragraphs[0].text = "Pie de página"

    buffer = io.BytesIO()
    documento.save(buffer)
    return buffer.getvalue()


def test_extrae_parrafos_con_su_indice_como_ubicacion() -> None:
    segmentos = leer_texto_docx(io.BytesIO(_docx_de_prueba()))

    parrafos = [s for s in segmentos if s.ubicacion.startswith("Párrafo")]
    assert any(s.texto == "Título del documento" for s in parrafos)
    assert any("Texto normal y " in s.texto for s in parrafos)


def test_extrae_celdas_de_tabla_con_fila_y_columna() -> None:
    segmentos = leer_texto_docx(io.BytesIO(_docx_de_prueba()))

    celdas = {s.ubicacion: s.texto for s in segmentos if s.ubicacion.startswith("Tabla")}
    assert celdas["Tabla 1, fila 1, columna 1"] == "Encabezado columna 1"
    assert celdas["Tabla 1, fila 2, columna 2"] == "Valor B"


def test_extrae_encabezado_y_pie_de_pagina() -> None:
    segmentos = leer_texto_docx(io.BytesIO(_docx_de_prueba()))

    textos = {s.texto for s in segmentos}
    assert "Encabezado de página" in textos
    assert "Pie de página" in textos


def test_no_extrae_parrafos_vacios() -> None:
    documento = Document()
    documento.add_paragraph("")
    documento.add_paragraph("   ")
    documento.add_paragraph("Contenido real")
    buffer = io.BytesIO()
    documento.save(buffer)

    segmentos = leer_texto_docx(io.BytesIO(buffer.getvalue()))

    assert len(segmentos) == 1
    assert segmentos[0].texto == "Contenido real"


def test_conserva_las_dos_corridas_del_parrafo_con_negrita_como_un_solo_segmento() -> None:
    """RF-14 (aplicar solo lo aceptado sin romper formato) necesita poder
    re-ubicar el párrafo por índice más adelante; acá solo se verifica que
    el texto del párrafo completo -- con negrita y sin ella -- llega junto,
    ubicación = índice de párrafo, no una por corrida (`run`)."""
    segmentos = leer_texto_docx(io.BytesIO(_docx_de_prueba()))

    con_negrita = [s for s in segmentos if "texto en negrita" in s.texto]
    assert len(con_negrita) == 1
    assert con_negrita[0].texto == "Texto normal y texto en negrita"


def test_aplicar_correcciones_corrige_parrafo_tabla_y_encabezado() -> None:
    correcciones = [
        CorreccionAplicable(
            "Párrafo 0", "Título del documento", "Título del documento correcto"
        ),
        CorreccionAplicable("Tabla 1, fila 2, columna 2", "Valor B", "Valor corregido"),
        CorreccionAplicable(
            "Encabezado 1, párrafo 0", "Encabezado de página", "Encabezado correcto"
        ),
        CorreccionAplicable("Pie de página 1, párrafo 0", "Pie de página", "Pie corregido"),
    ]

    resultado = aplicar_correcciones_docx(io.BytesIO(_docx_de_prueba()), correcciones)

    assert resultado.no_aplicadas == []
    segmentos = leer_texto_docx(io.BytesIO(resultado.contenido))
    textos = {s.texto for s in segmentos}
    assert "Título del documento correcto" in textos
    assert "Encabezado correcto" in textos
    assert "Pie corregido" in textos
    celdas = {s.ubicacion: s.texto for s in segmentos if s.ubicacion.startswith("Tabla")}
    assert celdas["Tabla 1, fila 2, columna 2"] == "Valor corregido"


def test_aplicar_correcciones_preserva_la_negrita() -> None:
    correcciones = [CorreccionAplicable("Párrafo 1", "texto en negrita", "texto en cursiva")]

    resultado = aplicar_correcciones_docx(io.BytesIO(_docx_de_prueba()), correcciones)

    documento = Document(io.BytesIO(resultado.contenido))
    parrafo = documento.paragraphs[1]
    corrida_en_negrita = next(c for c in parrafo.runs if c.bold)
    assert corrida_en_negrita.text == "texto en cursiva"


def test_correccion_con_texto_inexistente_queda_en_no_aplicadas() -> None:
    correccion = CorreccionAplicable("Párrafo 0", "texto que no está ahí", "reemplazo")

    resultado = aplicar_correcciones_docx(io.BytesIO(_docx_de_prueba()), [correccion])

    assert resultado.no_aplicadas == [correccion]


def test_construir_docx_desde_segmentos_aplica_correccion_y_conserva_los_demas() -> None:
    segmentos = [
        SegmentoTexto(
            texto="Se encontró un depositos sin registrar.", ubicacion="Página 1, bloque 1"
        ),
        SegmentoTexto(texto="Segundo bloque sin errores.", ubicacion="Página 1, bloque 2"),
    ]
    correcciones = [
        CorreccionAplicable("Página 1, bloque 1", "depositos", "depósitos"),
    ]

    contenido = construir_docx_desde_segmentos(segmentos, correcciones)

    documento = Document(io.BytesIO(contenido))
    texto_completo = "\n".join(p.text for p in documento.paragraphs)
    assert "depósitos" in texto_completo
    assert "depositos" not in texto_completo
    assert "Segundo bloque sin errores." in texto_completo
