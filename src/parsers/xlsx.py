"""Extractor de texto de Excel genérico (RF-10, CU-05): celdas de texto de
cualquier hoja, sin asumir un esquema de columnas -- distinto de
`parsers.excel` (CU-01), que sí asume el esquema contable Cuenta|Descripcion|
...|Moneda. Solo lectura -- la revisión ortográfica vive en `src/ortografia`.
"""

from __future__ import annotations

import io
import re

from openpyxl import load_workbook

from parsers.correcciones import CorreccionAplicable, ResultadoCorreccion
from parsers.runs import reemplazar_primera_ocurrencia
from parsers.segmentos import SegmentoTexto


def leer_texto_xlsx(archivo: io.BytesIO | str) -> list[SegmentoTexto]:
    """Extrae solo celdas de texto: excluye fórmulas (se leen del libro sin
    `data_only` para detectarlas por `data_type == "f"`, sin importar si el
    resultado cacheado de la fórmula es texto) y celdas numéricas."""
    if hasattr(archivo, "seek"):
        archivo.seek(0)
    libro_formulas = load_workbook(archivo, data_only=False)
    if hasattr(archivo, "seek"):
        archivo.seek(0)
    libro_valores = load_workbook(archivo, data_only=True)

    segmentos: list[SegmentoTexto] = []
    for nombre_hoja in libro_formulas.sheetnames:
        hoja_formulas = libro_formulas[nombre_hoja]
        hoja_valores = libro_valores[nombre_hoja]
        filas = zip(hoja_formulas.iter_rows(), hoja_valores.iter_rows(), strict=True)
        for fila_formulas, fila_valores in filas:
            celdas = zip(fila_formulas, fila_valores, strict=True)
            for celda_formula, celda_valor in celdas:
                if celda_formula.data_type == "f":
                    continue
                valor = celda_valor.value
                if isinstance(valor, str) and valor.strip():
                    segmentos.append(
                        SegmentoTexto(
                            texto=valor,
                            ubicacion=f"{nombre_hoja}!{celda_valor.coordinate}",
                        )
                    )
    return segmentos


_PATRON_UBICACION = re.compile(r"^(.+)!([A-Z]+\d+)$")


def aplicar_correcciones_xlsx(
    archivo: io.BytesIO | str, correcciones: list[CorreccionAplicable]
) -> ResultadoCorreccion:
    """Las celdas de texto no tienen corridas con formato propio que
    preservar (a diferencia de docx/pptx) -- un reemplazo de subcadena sobre
    el valor de la celda alcanza."""
    if hasattr(archivo, "seek"):
        archivo.seek(0)
    libro = load_workbook(archivo, data_only=False)

    no_aplicadas: list[CorreccionAplicable] = []
    for correccion in correcciones:
        coincidencia = _PATRON_UBICACION.match(correccion.ubicacion)
        aplicada = False
        if coincidencia:
            nombre_hoja, coordenada = coincidencia.groups()
            if nombre_hoja in libro.sheetnames:
                celda = libro[nombre_hoja][coordenada]
                if isinstance(celda.value, str):
                    nuevo_valor = reemplazar_primera_ocurrencia(
                        celda.value, correccion.texto_original, correccion.texto_nuevo
                    )
                    if nuevo_valor is not None:
                        celda.value = nuevo_valor
                        aplicada = True
        if not aplicada:
            no_aplicadas.append(correccion)

    buffer = io.BytesIO()
    libro.save(buffer)
    return ResultadoCorreccion(contenido=buffer.getvalue(), no_aplicadas=no_aplicadas)
