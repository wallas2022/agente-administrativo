"""Pruebas de exportación a .docx con palabras dudosas resaltadas (CU-06,
Bloque 3)."""

import io

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from ocr.exportar_docx import construir_docx_resaltado


def _parrafos(contenido_docx: bytes) -> list[str]:
    documento = Document(io.BytesIO(contenido_docx))
    return [p.text for p in documento.paragraphs]


def test_construir_docx_conserva_el_texto_exacto() -> None:
    texto = "Factura Q 1,250.00\nSegunda línea"
    contenido = construir_docx_resaltado(texto, nivel_por_palabra={})
    assert _parrafos(contenido) == ["Factura Q 1,250.00", "Segunda línea"]


def test_construir_docx_resalta_palabra_dudosa_en_rojo() -> None:
    contenido = construir_docx_resaltado(
        "Hola mllndo", nivel_por_palabra={"mllndo": "dudosa"}
    )
    documento = Document(io.BytesIO(contenido))
    runs = documento.paragraphs[0].runs
    coincidencias = [r for r in runs if r.text == "mllndo"]
    assert len(coincidencias) == 1
    assert coincidencias[0].font.highlight_color == WD_COLOR_INDEX.RED


def test_construir_docx_resalta_palabra_a_revisar_en_amarillo() -> None:
    contenido = construir_docx_resaltado("Hola Alg0", nivel_por_palabra={"Alg0": "revisar"})
    documento = Document(io.BytesIO(contenido))
    runs = documento.paragraphs[0].runs
    coincidencias = [r for r in runs if r.text == "Alg0"]
    assert len(coincidencias) == 1
    assert coincidencias[0].font.highlight_color == WD_COLOR_INDEX.YELLOW


def test_construir_docx_no_resalta_palabras_sin_nivel() -> None:
    contenido = construir_docx_resaltado("Hola Mundo", nivel_por_palabra={"Mundo": "dudosa"})
    documento = Document(io.BytesIO(contenido))
    runs = documento.paragraphs[0].runs
    hola = next(r for r in runs if r.text == "Hola")
    assert hola.font.highlight_color is None
