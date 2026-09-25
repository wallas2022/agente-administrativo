"""Salida de CU-01: Excel con celdas marcadas + JSON de hallazgos (RF-12, RF-14,
RF-15). No aplica ninguna corrección automática de datos — solo resalta
visualmente lo que ya detectó `reglas.py`; aplicar o no un hallazgo lo decide
el Revisor (CU-07).
"""

from __future__ import annotations

import dataclasses
import io
import json

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from validadores.contable.reglas import HallazgoDetectado

_COLOR_POR_SEVERIDAD = {
    "alta": "FFFF0000",  # rojo
    "media": "FFFFA500",  # naranja
    "baja": "FFFFFF00",  # amarillo
}


def marcar_celdas_en_libro(
    archivo_original: io.BytesIO, hallazgos: list[HallazgoDetectado]
) -> bytes:
    archivo_original.seek(0)
    wb = load_workbook(io.BytesIO(archivo_original.read()))

    for h in hallazgos:
        celda = wb[h.hoja].cell(row=h.fila, column=1)
        color = _COLOR_POR_SEVERIDAD.get(h.severidad, "FFFFFF00")
        celda.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        texto_comentario = f"[{h.regla_codigo}] {h.descripcion}"
        if celda.comment is not None:
            texto_comentario = f"{celda.comment.text}\n{texto_comentario}"
        celda.comment = Comment(texto_comentario, "Agente Administrativo")

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def hallazgos_a_json(hallazgos: list[HallazgoDetectado]) -> str:
    return json.dumps([dataclasses.asdict(h) for h in hallazgos], ensure_ascii=False, indent=2)
