"""Pruebas de integración del pipeline CU-01 dentro del orquestador
(Sprint 1, item 6): parser → reglas → RAG → explicación → salida, disparado
por `ejecutar_analisis` cuando `tipo_revision == "contable"`.
"""

import io
import json
import uuid
from datetime import UTC, date, datetime, timedelta

import boto3
from moto import mock_aws
from openpyxl import Workbook, load_workbook
from qdrant_client import QdrantClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from comun.estados import EstadoAnalisis, EstadoDocumento
from comun.modelos import (
    Analisis,
    Area,
    Base,
    Documento,
    Hallazgo,
    Rol,
    TipoRevision,
    Usuario,
    VersionDocumento,
)
from orquestador.pipeline_contable import COLECCION_RAG
from orquestador.tareas import ejecutar_analisis
from rag.ingesta import ingerir_fragmentos

BUCKET = "documentos"


def _libro_con_cuenta_inexistente() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Partidas"
    ws.append(["Cuenta", "Descripcion", "Fecha", "Debe", "Haber", "Moneda"])
    ws.append(["1010", "Cobro a cliente", date(2026, 1, 5), 100.0, 0.0, "Q"])
    ws.append(["9999", "Cuenta que no existe", date(2026, 1, 5), 0.0, 100.0, "Q"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _libro_con_descuadre_y_cuenta_inexistente() -> bytes:
    """RN-02 (fila 3, plantilla) + RN-01 (fila 4, la última partida — LLM)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Partidas"
    ws.append(["Cuenta", "Descripcion", "Fecha", "Debe", "Haber", "Moneda"])
    ws.append(["1010", "Cobro a cliente", date(2026, 1, 5), 100.0, 0.0, "Q"])
    ws.append(["9999", "Cuenta que no existe", date(2026, 1, 5), 0.0, 30.0, "Q"])
    ws.append(["4010", "Ingreso por servicio", date(2026, 1, 5), 0.0, 40.0, "Q"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _libro_sin_errores() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Partidas"
    ws.append(["Cuenta", "Descripcion", "Fecha", "Debe", "Haber", "Moneda"])
    ws.append(["1010", "Cobro a cliente", date(2026, 1, 5), 100.0, 0.0, "Q"])
    ws.append(["4010", "Ingreso por servicio", date(2026, 1, 5), 0.0, 100.0, "Q"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _sesion_con_documento_analisis_y_version(
    contenido: bytes,
    *,
    periodo_cierre: str = "2026-01",
) -> tuple[Session, uuid.UUID, uuid.UUID, str]:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    sesion = Session(engine)

    area = Area(nombre="Contabilidad")
    rol = Rol(nombre="analista")
    sesion.add_all([area, rol])
    sesion.flush()

    usuario = Usuario(nombre="Ana Lista", email="ana@ejemplo.gt", area_id=area.id, rol_id=rol.id)
    sesion.add(usuario)
    sesion.flush()

    documento = Documento(
        nombre_original="cierre.xlsx",
        tipo_archivo="xlsx",
        tamano_bytes=len(contenido),
        area_id=area.id,
        usuario_carga_id=usuario.id,
        fecha_carga=datetime.now(UTC),
        fecha_expiracion=date.today() + timedelta(days=90),
        estado=EstadoDocumento.CARGADO.value,
    )
    tipo_revision = TipoRevision(nombre="contable")
    sesion.add_all([documento, tipo_revision])
    sesion.flush()

    analisis = Analisis(
        documento_id=documento.id,
        tipo_revision_id=tipo_revision.id,
        usuario_id=usuario.id,
        fecha_inicio=datetime.now(UTC),
        estado=EstadoAnalisis.PROCESANDO.value,
        periodo_cierre=periodo_cierre,
    )
    sesion.add(analisis)
    sesion.flush()

    llave = f"{area.id}/{documento.id}/cierre.xlsx"
    version = VersionDocumento(
        documento_id=documento.id,
        numero_version=1,
        ruta_almacenamiento=llave,
        es_corregida=False,
        fecha_creacion=datetime.now(UTC),
    )
    sesion.add(version)
    sesion.commit()

    return sesion, documento.id, analisis.id, llave


def _embedding_falso(_texto: str) -> list[float]:
    return [0.0] * 4


def _llm_falso(_prompt: str) -> str:
    return "Causa probable: prueba. Corrección sugerida: prueba."


def test_ejecutar_analisis_contable_persiste_hallazgos_y_marca_con_hallazgos() -> None:
    contenido = _libro_con_cuenta_inexistente()
    sesion, documento_id, analisis_id, llave = _sesion_con_documento_analisis_y_version(contenido)

    with mock_aws():
        cliente_s3 = boto3.client("s3", region_name="us-east-1")
        cliente_s3.create_bucket(Bucket=BUCKET)
        cliente_s3.put_object(Bucket=BUCKET, Key=llave, Body=contenido)

        resultado = ejecutar_analisis(
            sesion,
            documento_id,
            analisis_id,
            cliente_s3=cliente_s3,
            bucket=BUCKET,
            cliente_qdrant=QdrantClient(":memory:"),
            funcion_embedding=_embedding_falso,
            funcion_llm=_llm_falso,
            modelo_llm="modelo-de-prueba",
        )

        assert resultado == EstadoDocumento.CON_HALLAZGOS.value

        documento = sesion.get(Documento, documento_id)
        analisis = sesion.get(Analisis, analisis_id)
        assert documento.estado == EstadoDocumento.CON_HALLAZGOS.value
        assert analisis.estado == EstadoAnalisis.COMPLETADO.value
        # RN-02 se explica por plantilla (sin LLM) — no debió llamarse a _llm_falso.
        assert analisis.modelo_llm is None
        assert analisis.version_prompt is None
        assert float(analisis.total_debe) == 100.0
        assert float(analisis.total_haber) == 100.0
        assert analisis.moneda == "Q"

        hallazgos = sesion.query(Hallazgo).filter_by(analisis_id=analisis_id).all()
        assert len(hallazgos) == 1
        assert hallazgos[0].ubicacion == "Partidas!A3"
        assert hallazgos[0].severidad == "alta"
        assert "catálogo de cuentas" in hallazgos[0].correccion_sugerida
        assert hallazgos[0].fuente_citada is None  # sin fragmentos ingeridos en este Qdrant

        versiones = (
            sesion.query(VersionDocumento)
            .filter_by(documento_id=documento_id)
            .order_by(VersionDocumento.numero_version)
            .all()
        )
        assert len(versiones) == 2
        assert versiones[1].es_corregida is True

        llave_marcada = versiones[1].ruta_almacenamiento
        libro_marcado_bytes = cliente_s3.get_object(Bucket=BUCKET, Key=llave_marcada)["Body"].read()
        wb_marcado = load_workbook(io.BytesIO(libro_marcado_bytes))
        assert wb_marcado["Partidas"]["A3"].comment is not None


def test_ejecutar_analisis_contable_sin_hallazgos_marca_en_revision() -> None:
    contenido = _libro_sin_errores()
    sesion, documento_id, analisis_id, llave = _sesion_con_documento_analisis_y_version(contenido)

    with mock_aws():
        cliente_s3 = boto3.client("s3", region_name="us-east-1")
        cliente_s3.create_bucket(Bucket=BUCKET)
        cliente_s3.put_object(Bucket=BUCKET, Key=llave, Body=contenido)

        resultado = ejecutar_analisis(
            sesion,
            documento_id,
            analisis_id,
            cliente_s3=cliente_s3,
            bucket=BUCKET,
            cliente_qdrant=QdrantClient(":memory:"),
            funcion_embedding=_embedding_falso,
            funcion_llm=_llm_falso,
            modelo_llm="modelo-de-prueba",
        )

        assert resultado == EstadoDocumento.EN_REVISION.value
        assert sesion.query(Hallazgo).filter_by(analisis_id=analisis_id).count() == 0

        versiones = sesion.query(VersionDocumento).filter_by(documento_id=documento_id).all()
        assert len(versiones) == 1  # no se sube versión corregida si no hay hallazgos

        # Los totales se calculan siempre, haya o no hallazgos (Pantalla 3, U4).
        analisis = sesion.get(Analisis, analisis_id)
        assert float(analisis.total_debe) == 100.0
        assert float(analisis.total_haber) == 100.0
        assert analisis.moneda == "Q"


def test_rn03_usa_periodo_cierre_del_analisis_no_la_fecha_en_que_corre() -> None:
    """El período que se está cerrando lo declara quien sube el documento
    (`periodo_cierre`), no depende de cuándo corre el worker: un documento de
    enero revisado hoy (2026-09) no debe disparar RN-03 si periodo_cierre
    coincide con enero; sí debe dispararlo si se declara un período distinto.
    """
    contenido = _libro_sin_errores()  # partidas con fecha 2026-01-05, sin otros errores

    sesion, documento_id, analisis_id, llave = _sesion_con_documento_analisis_y_version(
        contenido, periodo_cierre="2026-01"
    )
    with mock_aws():
        cliente_s3 = boto3.client("s3", region_name="us-east-1")
        cliente_s3.create_bucket(Bucket=BUCKET)
        cliente_s3.put_object(Bucket=BUCKET, Key=llave, Body=contenido)

        ejecutar_analisis(
            sesion,
            documento_id,
            analisis_id,
            cliente_s3=cliente_s3,
            bucket=BUCKET,
            cliente_qdrant=QdrantClient(":memory:"),
            funcion_embedding=_embedding_falso,
            funcion_llm=_llm_falso,
            modelo_llm="modelo-de-prueba",
        )
        assert sesion.query(Hallazgo).filter_by(analisis_id=analisis_id).count() == 0

    sesion2, documento_id2, analisis_id2, llave2 = _sesion_con_documento_analisis_y_version(
        contenido, periodo_cierre="2026-02"
    )
    with mock_aws():
        cliente_s3 = boto3.client("s3", region_name="us-east-1")
        cliente_s3.create_bucket(Bucket=BUCKET)
        cliente_s3.put_object(Bucket=BUCKET, Key=llave2, Body=contenido)

        ejecutar_analisis(
            sesion2,
            documento_id2,
            analisis_id2,
            cliente_s3=cliente_s3,
            bucket=BUCKET,
            cliente_qdrant=QdrantClient(":memory:"),
            funcion_embedding=_embedding_falso,
            funcion_llm=_llm_falso,
            modelo_llm="modelo-de-prueba",
        )
        hallazgos = sesion2.query(Hallazgo).filter_by(analisis_id=analisis_id2).all()
        assert {h.ubicacion for h in hallazgos} == {"Partidas!A2", "Partidas!A3"}
        assert all("período" in h.descripcion for h in hallazgos)


def test_hallazgo_guarda_la_fuente_citada_por_rag() -> None:
    """Pantalla 3 (U4): cada hallazgo debe mostrar la fuente citada (RF-12,
    PP-07) — se persiste en Hallazgo.fuente_citada, no solo queda en el
    texto libre de correccion_sugerida."""
    contenido = _libro_con_cuenta_inexistente()
    sesion, documento_id, analisis_id, llave = _sesion_con_documento_analisis_y_version(contenido)

    cliente_qdrant = QdrantClient(":memory:")
    ingerir_fragmentos(
        cliente_qdrant,
        coleccion=COLECCION_RAG,
        fuente_id="politica-cierre-contable",
        fragmentos=[("sec-2-cuentas", "Toda cuenta usada debe existir en el catálogo vigente.")],
        funcion_embedding=_embedding_falso,
        dimension=4,
    )

    with mock_aws():
        cliente_s3 = boto3.client("s3", region_name="us-east-1")
        cliente_s3.create_bucket(Bucket=BUCKET)
        cliente_s3.put_object(Bucket=BUCKET, Key=llave, Body=contenido)

        ejecutar_analisis(
            sesion,
            documento_id,
            analisis_id,
            cliente_s3=cliente_s3,
            bucket=BUCKET,
            cliente_qdrant=cliente_qdrant,
            funcion_embedding=_embedding_falso,
            funcion_llm=_llm_falso,
            modelo_llm="modelo-de-prueba",
        )

        hallazgo = sesion.query(Hallazgo).filter_by(analisis_id=analisis_id).one()
        assert hallazgo.fuente_citada == "politica-cierre-contable"


def test_pipeline_agrupa_llm_en_una_sola_llamada_y_persiste_deterministas_primero() -> None:
    """RNF-04 (ver docs/04-pruebas/resultados/local-S1.md): con documentos de
    varias decenas de hallazgos, una llamada al LLM por hallazgo tardaba
    20-40 min. Esta prueba fija ese comportamiento: RN-01 y RN-05 se agrupan
    en una sola llamada, y los hallazgos por plantilla (RN-02 acá) quedan
    persistidos y consultables ANTES de que esa llamada termine."""
    contenido = _libro_con_descuadre_y_cuenta_inexistente()
    sesion, documento_id, analisis_id, llave = _sesion_con_documento_analisis_y_version(contenido)

    llamadas_llm: list[str] = []
    ubicaciones_visibles_durante_la_llamada_llm: list[str] = []

    def llm_falso_que_verifica_orden(prompt: str) -> str:
        llamadas_llm.append(prompt)
        # El hallazgo determinista (RN-02) debe estar ya persistido y
        # consultable en este punto -- antes de que la llamada al LLM
        # siquiera termine de "resolver".
        existentes = sesion.query(Hallazgo).filter_by(analisis_id=analisis_id).all()
        ubicaciones_visibles_durante_la_llamada_llm.extend(h.ubicacion for h in existentes)
        return json.dumps(
            [{"indice": 0, "causa_probable": "causa llm", "correccion_sugerida": "correccion llm"}]
        )

    with mock_aws():
        cliente_s3 = boto3.client("s3", region_name="us-east-1")
        cliente_s3.create_bucket(Bucket=BUCKET)
        cliente_s3.put_object(Bucket=BUCKET, Key=llave, Body=contenido)

        ejecutar_analisis(
            sesion,
            documento_id,
            analisis_id,
            cliente_s3=cliente_s3,
            bucket=BUCKET,
            cliente_qdrant=QdrantClient(":memory:"),
            funcion_embedding=_embedding_falso,
            funcion_llm=llm_falso_que_verifica_orden,
            modelo_llm="gpt-oss:20b",
        )

        assert len(llamadas_llm) == 1  # una sola llamada, sin importar cuántos la necesiten
        assert ubicaciones_visibles_durante_la_llamada_llm == ["Partidas!A3"]  # RN-02 ya estaba

        hallazgos = sesion.query(Hallazgo).filter_by(analisis_id=analisis_id).all()
        assert len(hallazgos) == 2
        por_ubicacion = {h.ubicacion: h for h in hallazgos}
        assert "catálogo de cuentas" in por_ubicacion["Partidas!A3"].correccion_sugerida
        assert "causa llm" in por_ubicacion["Partidas!A4"].correccion_sugerida
        assert "correccion llm" in por_ubicacion["Partidas!A4"].correccion_sugerida


def test_ejecutar_analisis_sin_dependencias_de_infraestructura_usa_flujo_generico() -> None:
    """Si no se inyectan las dependencias de RAG/S3 (p. ej. un tipo de revisión
    sin adaptador todavía), se preserva el comportamiento genérico de L2."""
    contenido = _libro_con_cuenta_inexistente()
    sesion, documento_id, analisis_id, _llave = _sesion_con_documento_analisis_y_version(contenido)

    resultado = ejecutar_analisis(sesion, documento_id, analisis_id)

    assert resultado == EstadoDocumento.EN_REVISION.value
    assert sesion.query(Hallazgo).filter_by(analisis_id=analisis_id).count() == 0
